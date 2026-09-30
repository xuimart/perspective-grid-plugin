from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.space_before = Pt(13 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(7)

def inline(p, text):
    pattern = r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^\s]+\))'
    for token in re.split(pattern, text):
        if token.startswith('**') and token.endswith('**'):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith('`') and token.endswith('`'):
            r = p.add_run(token[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(10)
        elif re.match(r'\[.*\]\(https?://', token):
            label, url = re.match(r'\[(.*?)\]\((.*)\)', token).groups()
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '245D77'); prop.append(color)
            r.append(prop); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)
        else:
            p.add_run(token)

lines = (ROOT / 'PROMPT-COMPLETO-PERSPECTIVE-GRID-UXP.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1; continue
    if line.startswith('```'):
        block = []; i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            block.append(lines[i]); i += 1
        for n, code in enumerate(block):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = n < len(block) - 1
            r = p.add_run(code); r.font.name = 'Consolas'; r.font.size = Pt(9)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [s.strip() for s in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', s) for s in cells): rows.append(cells)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0])); table.autofit = False
        widths = [1.5, 2.8, 2.7] if len(rows[0]) == 3 else [7 / len(rows[0])] * len(rows[0])
        for col, width in zip(table.columns, widths): col.width = Inches(width)
        pr = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
        pr.append(borders)
        for row_index, data in enumerate(rows):
            row = table.add_row()
            rp = row._tr.get_or_add_trPr()
            rp.append(OxmlElement('w:cantSplit'))
            if row_index == 0: rp.append(OxmlElement('w:tblHeader'))
            for cell, value, width in zip(row.cells, data, widths):
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
                for edge in ['top', 'left', 'bottom', 'right']:
                    e = OxmlElement('w:' + edge); e.set(qn('w:w'), '90'); e.set(qn('w:type'), 'dxa'); margins.append(e)
                cp.append(margins)
                if row_index == 0:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8ECEE'); cp.append(shade)
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.02
                inline(p, value)
                for r in p.runs: r.font.size = Pt(10); r.bold = row_index == 0 or r.bold
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        continue
    heading = re.match(r'^(#{1,4}) (.*)', line)
    if heading:
        level = len(heading.group(1))
        p = doc.add_paragraph(style='Title' if level == 1 else 'Heading ' + str(level - 1))
        inline(p, heading.group(2))
    elif re.match(r'^(\d+\. |[-*] )', line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.17)
        p.paragraph_format.first_line_indent = Inches(-.17)
        p.paragraph_format.space_after = Pt(4)
        inline(p, line)
    else:
        parts = [line]
        while i + 1 < len(lines) and lines[i+1].strip() and not re.match(r'^(#|\||```|\d+\. |[-*] )', lines[i+1]):
            i += 1; parts.append(lines[i])
        p = doc.add_paragraph(); inline(p, ' '.join(parts))
    i += 1

footer = section.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('Perspective Grid UXP  |  '); r.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Perspective Grid para Photoshop UXP'
doc.core_properties.subject = 'Especificacao funcional e tecnica e prompt de implementacao'
output = ROOT / 'Perspective-Grid-Especificacao-Completa-UXP.docx'
doc.save(output)
print(output)
